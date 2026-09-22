"""TXT / DOCX 文本提取与规范化。"""

import hashlib
import io
import re

import chardet
from docx import Document


class ExtractionError(Exception):
    """文本提取阶段错误，携带可定位的错误码。"""

    def __init__(self, code, message):
        super().__init__(message)
        self.code = code
        self.message = message


def decode_txt(raw: bytes) -> str:
    if not raw:
        raise ExtractionError("EMPTY_FILE", "文件内容为空。")
    detected = chardet.detect(raw[:65536]) or {}
    encoding = detected.get("encoding") or "utf-8"
    try:
        text = raw.decode(encoding, errors="strict")
    except (UnicodeDecodeError, LookupError):
        try:
            text = raw.decode("utf-8", errors="replace")
        except Exception as exc:  # pragma: no cover - 理论上 replace 不会失败
            raise ExtractionError("DECODE_FAILED", f"文本解码失败：{exc}") from exc
    if not text.strip():
        raise ExtractionError("EMPTY_TEXT", "未从文件中提取到有效文本。")
    return text


def read_docx(raw: bytes) -> str:
    if not raw:
        raise ExtractionError("EMPTY_FILE", "文件内容为空。")
    try:
        doc = Document(io.BytesIO(raw))
    except Exception as exc:
        raise ExtractionError("DOCX_PARSE_FAILED", f"DOCX 解析失败：{exc}") from exc
    parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    # 表格中的文字也纳入（常见于带封面表格的作业）
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                t = cell.text.strip()
                if t:
                    parts.append(t)
    text = "\n\n".join(parts)
    if not text.strip():
        raise ExtractionError("EMPTY_TEXT", "DOCX 中没有可提取的段落文本。")
    return text


def extract(source_format: str, raw: bytes) -> str:
    if source_format == "txt":
        text = decode_txt(raw)
    elif source_format == "docx":
        text = read_docx(raw)
    else:  # pragma: no cover - 上传层已限制
        raise ExtractionError("UNSUPPORTED_FORMAT", f"不支持的格式：{source_format}")
    return normalize_text(text)


def normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()


def content_hash_of(text: str) -> str:
    """规范化文本的 SHA-256，作为去重与结果绑定用的内容摘要。"""
    return hashlib.sha256(normalize_text(text).encode("utf-8")).hexdigest()
